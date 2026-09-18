"""Build the synthetic Phase 3 fixture documents programmatically (no binary files are committed).

    python -m tests.fixtures.documents.builders /some/output/dir     # write them for manual inspection
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from tests.fixtures.documents.synthetic_content import CONTRACT_AR, CONTRACT_EN, CV_AR, CV_EN


def load_pymupdf():
    """PyMuPDF >= 1.24.3 is importable as "pymupdf"; older versions (hr_assistant pins 1.23.8) only as "fitz"."""
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    return pymupdf


FIXTURE_NAMES = (
    "sample_contract_ar.docx",
    "sample_contract_en.docx",
    "sample_cv_ar.docx",
    "sample_cv_en.docx",
    "sample_contract_text.pdf",
    "sample_empty.pdf",
    "sample_scanned.pdf",
)


# --------------------------------------------------------------------------- DOCX
def _rtl(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    properties.append(bidi)
    return paragraph


def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def contract_docx(content: dict, *, rtl: bool = False) -> bytes:
    from docx import Document

    fix = _rtl if rtl else (lambda paragraph: paragraph)
    document = Document()
    fix(document.add_heading(content["title"], level=0))
    fix(document.add_paragraph(content["intro"]))
    for label, value in content["fields"]:
        fix(document.add_paragraph(f"{label}: {value}"))
    fix(document.add_heading(content["compensation_heading"], level=1))
    table = document.add_table(rows=0, cols=2)
    for label, value in content["compensation_rows"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    for heading, text in content["clauses"]:
        fix(document.add_heading(heading, level=1))
        for line in text.split("\n"):
            fix(document.add_paragraph(line))
    return _docx_bytes(document)


def cv_docx(content: dict, *, rtl: bool = False) -> bytes:
    from docx import Document

    fix = _rtl if rtl else (lambda paragraph: paragraph)
    document = Document()
    fix(document.add_heading(content["name"], level=0))
    fix(document.add_paragraph(content["contact"]))
    if "phone" in content:
        fix(document.add_paragraph(content["phone"]))
    fix(document.add_paragraph(content["location"]))
    fix(document.add_heading(content["summary_heading"], level=1))
    fix(document.add_paragraph(content["summary"]))
    fix(document.add_heading(content["skills_heading"], level=1))
    if "skills" in content:
        for skill in content["skills"]:
            fix(document.add_paragraph(skill, style="List Bullet"))
    else:
        fix(document.add_paragraph(content["skills_line"]))
    fix(document.add_heading(content["experience_heading"], level=1))
    for header, dates, bullets in content["experience"]:
        fix(document.add_paragraph(header))
        fix(document.add_paragraph(dates))
        for bullet in bullets:
            fix(document.add_paragraph(bullet, style="List Bullet"))
    for key in ("education", "certifications", "languages"):
        fix(document.add_heading(content[f"{key}_heading"], level=1))
        for line in content[key]:
            fix(document.add_paragraph(line))
    return _docx_bytes(document)


def docx_from_paragraphs(paragraphs: list[str], *, table_rows: list[tuple[str, ...]] | None = None) -> bytes:
    """Small ad-hoc documents for focused tests."""
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
    return _docx_bytes(document)


def empty_docx() -> bytes:
    from docx import Document

    return _docx_bytes(Document())


# --------------------------------------------------------------------------- PDF
def contract_text_pdf() -> bytes:
    pymupdf = load_pymupdf()

    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    y = 60
    page.insert_text((50, y), CONTRACT_EN["title"], fontsize=18, fontname="hebo")
    y += 30
    page.insert_text((50, y), CONTRACT_EN["intro"], fontsize=10, fontname="helv")
    y += 22
    for label, value in CONTRACT_EN["fields"]:
        page.insert_text((50, y), f"{label}: {value}", fontsize=10, fontname="helv")
        y += 16
    y += 14
    page.insert_text((50, y), CONTRACT_EN["compensation_heading"], fontsize=14, fontname="hebo")
    y += 12
    top, row_height, rows = y, 20, CONTRACT_EN["compensation_rows"]
    left, middle, right = 50, 250, 500
    for index in range(len(rows) + 1):
        page.draw_line((left, top + index * row_height), (right, top + index * row_height))
    for x in (left, middle, right):
        page.draw_line((x, top), (x, top + len(rows) * row_height))
    for index, (label, value) in enumerate(rows):
        baseline = top + index * row_height + 14
        page.insert_text((left + 5, baseline), label, fontsize=10, fontname="helv")
        page.insert_text((middle + 5, baseline), value, fontsize=10, fontname="helv")
    y = top + len(rows) * row_height + 30
    for heading, text in CONTRACT_EN["clauses"][:3]:
        page.insert_text((50, y), heading, fontsize=14, fontname="hebo")
        y += 18
        page.insert_text((50, y), text, fontsize=10, fontname="helv")
        y += 26

    second = document.new_page(width=595, height=842)
    y = 60
    for heading, text in CONTRACT_EN["clauses"][3:]:
        second.insert_text((50, y), heading, fontsize=14, fontname="hebo")
        y += 18
        for line in _wrap(text, 95):
            second.insert_text((50, y), line, fontsize=10, fontname="helv")
            y += 14
        y += 12
    return document.tobytes()


def _wrap(text: str, width: int) -> list[str]:
    words, lines, current = text.split(), [], ""
    for word in words:
        if len(current) + len(word) + 1 > width:
            lines.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    return [*lines, current] if current else lines


def empty_pdf() -> bytes:
    pymupdf = load_pymupdf()

    document = pymupdf.open()
    document.new_page(width=595, height=842)
    return document.tobytes()


def scanned_pdf() -> bytes:
    """A page that only contains a picture of text: no text layer."""
    pymupdf = load_pymupdf()

    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(50, 50, 545, 250), stream=image_bytes("PNG"))
    return document.tobytes()


def pdf_with_text_pages_and_scanned_page() -> bytes:
    pymupdf = load_pymupdf()

    document = pymupdf.open(stream=contract_text_pdf(), filetype="pdf")
    scanned = pymupdf.open(stream=scanned_pdf(), filetype="pdf")
    document.insert_pdf(scanned)
    return document.tobytes()


# --------------------------------------------------------------------------- images
def image_bytes(fmt: str = "PNG", size: tuple[int, int] = (480, 160)) -> bytes:
    """A picture of text rendered with PyMuPDF (no Pillow): used inside scanned PDFs and for rejection tests."""
    pymupdf = load_pymupdf()
    document = pymupdf.open()
    page = document.new_page(width=size[0], height=size[1])
    page.insert_text((20, 50), "EMPLOYMENT CONTRACT", fontsize=14)
    page.insert_text((20, 90), "Basic Salary: 12,000 SAR", fontsize=12)
    pixmap = page.get_pixmap(dpi=72, alpha=False)
    return pixmap.tobytes("jpg" if fmt.upper() in ("JPG", "JPEG") else "png")


# --------------------------------------------------------------------------- all
def build_all() -> dict[str, bytes]:
    return {
        "sample_contract_ar.docx": contract_docx(CONTRACT_AR, rtl=True),
        "sample_contract_en.docx": contract_docx(CONTRACT_EN),
        "sample_cv_ar.docx": cv_docx(CV_AR, rtl=True),
        "sample_cv_en.docx": cv_docx(CV_EN),
        "sample_contract_text.pdf": contract_text_pdf(),
        "sample_empty.pdf": empty_pdf(),
        "sample_scanned.pdf": scanned_pdf(),
    }


def write_all(directory: Path) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    paths = {}
    for name, data in build_all().items():
        path = directory / name
        path.write_bytes(data)
        paths[name] = path
    return paths


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: python -m tests.fixtures.documents.builders <output-dir>")
    for written in write_all(Path(sys.argv[1])).values():
        print(written)
