"""CV fields: only what the CV states, in the sections where it states it."""

from __future__ import annotations

import io

from docx import Document

from sanad.extraction import extract_cv
from sanad.models.common import ResultStatus
from sanad.models.extraction import CvExtraction, ExtractionMethodName, FieldStatus
from tests.fixtures.documents.builders import load_pymupdf
from tests.fixtures.documents.synthetic_content import CV_AR, CV_EN

pymupdf = load_pymupdf()
FOUND, NOT_FOUND, AMBIGUOUS = FieldStatus.FOUND, FieldStatus.NOT_FOUND, FieldStatus.AMBIGUOUS


def _cv(processor, build):
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = processor.parse_bytes(buffer.getvalue(), "cv.docx")
    return extract_cv(parsed), parsed


def test_english_cv_contact_details(parsed):
    result = extract_cv(parsed["sample_cv_en.docx"])
    assert result.status is ResultStatus.SUCCESS and result.document_type == "cv"
    assert result.name.value == "Jordan Sample" and result.name.confidence == "medium"
    assert result.name.sources[0].method is ExtractionMethodName.DOCUMENT_TITLE
    assert result.email.value == "jordan.sample@example.com"
    assert result.phone.value == "+966 55 000 1234"
    assert {s.method for s in result.phone.sources} == {ExtractionMethodName.LABELED_LINE, ExtractionMethodName.PATTERN}
    assert result.location.value == "Riyadh, Saudi Arabia"
    assert result.summary.value == CV_EN["summary"]


def test_english_cv_skills_education_certifications_languages(parsed):
    result = extract_cv(parsed["sample_cv_en.docx"])
    assert [item.text for item in result.skills.value] == CV_EN["skills"]
    [education] = result.education.value
    assert (education.degree, education.institution, education.start_year, education.end_year) == (
        "Bachelor of Science in Computer Science", "Example University", 2014, 2018)
    [certification] = result.certifications.value
    assert (certification.name, certification.year) == ("Certified Data Professional", 2022)
    assert [(l.language, l.level) for l in result.languages.value] == [("Arabic", "Native"), ("English", "Fluent")]


def test_english_cv_work_experience_entries(parsed):
    result = extract_cv(parsed["sample_cv_en.docx"])
    first, second = result.work_experience.value
    assert (first.title, first.organization, first.start, first.end, first.is_current) == (
        "Data Analyst", "Example Analytics Co.", "Jan 2021", "Present", True)
    assert first.details == ["Built monthly sales dashboards.", "Automated data quality checks."]
    assert (second.title, second.organization, second.start, second.end, second.is_current) == (
        "Junior Analyst", "Sample Retail Group", "2018", "2020", False)
    assert not hasattr(result, "years_of_experience")  # experience is never computed in Phase 3


def test_arabic_cv(parsed):
    document = parsed["sample_cv_ar.docx"]
    result = extract_cv(document)
    assert result.name.value == CV_AR["name"]
    assert result.email.value == "reem.test@example.com"
    assert result.phone.value == "0550001234"
    assert result.location.value == "جدة"
    assert [s.text for s in result.skills.value] == ["إعداد القوائم المالية", "تحليل التكاليف", "Excel"]
    [job] = result.work_experience.value
    assert (job.title, job.organization, job.start, job.end, job.is_current) == (
        "محاسبة", "شركة مثال للتجارة", "2019", "حتى الآن", True)
    [education] = result.education.value
    assert (education.degree, education.institution, education.start_year, education.end_year) == (
        "بكالوريوس محاسبة", "جامعة المثال", 2015, 2019)
    assert [(l.language, l.level) for l in result.languages.value] == [("العربية", "اللغة الأم"), ("الإنجليزية", "متقدم")]
    for field in result.fields().values():
        for source in field.sources:
            assert source.text in document.full_text
        for item in field.value if isinstance(field.value, list) else []:
            assert item.source.text in document.full_text


def test_missing_cv_fields_are_not_found(processor):
    result, _ = _cv(processor, lambda d: (d.add_heading("Casey Placeholder", 0), d.add_paragraph("Email: casey@example.com")))
    assert result.name.status is FOUND and result.email.status is FOUND
    for name in ("phone", "location", "summary", "skills", "education", "certifications", "work_experience", "languages"):
        field = getattr(result, name)
        assert field.status is NOT_FOUND and field.value is None, name


def test_skills_and_experience_are_not_inferred_from_job_descriptions(processor):
    def build(d):
        d.add_heading("Casey Placeholder", 0)
        d.add_heading("Experience", 1)
        d.add_paragraph("Analyst, Example Co.")
        d.add_paragraph("Used Python and SQL daily for five years.")
    result, _ = _cv(processor, build)
    assert result.skills.status is NOT_FOUND
    [entry] = result.work_experience.value
    assert entry.title is None and entry.organization is None  # no explicit "X at Y" form
    assert entry.start is None and entry.header == "Analyst, Example Co."
    assert entry.details == ["Used Python and SQL daily for five years."]


def test_several_different_emails_are_ambiguous(processor):
    result, _ = _cv(processor, lambda d: (d.add_paragraph("Email: first@example.com"),
                                          d.add_paragraph("Alternative: second@example.com")))
    assert result.email.status is AMBIGUOUS and {c.value for c in result.email.candidates} == {
        "first@example.com", "second@example.com"}


def test_invalid_labeled_phone_is_ambiguous_not_guessed(processor):
    result, _ = _cv(processor, lambda d: d.add_paragraph("Phone: available on request"))
    assert result.phone.status is AMBIGUOUS and result.phone.value is None


def test_unrecognised_section_ends_the_previous_section(processor):
    def build(d):
        d.add_heading("Languages", 1)
        d.add_paragraph("English (Fluent)")
        d.add_heading("References", 1)
        d.add_paragraph("Available on request")
    result, _ = _cv(processor, build)
    assert [l.language for l in result.languages.value] == ["English"]


def test_cv_from_pdf_keeps_page_numbers(processor):
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((50, 60), "Casey Placeholder", fontsize=20, fontname="hebo")
    page.insert_text((50, 90), "Email: casey@example.com", fontsize=10)
    page.insert_text((50, 120), "Skills", fontsize=14, fontname="hebo")
    page.insert_text((50, 140), "Python, SQL, Tableau", fontsize=10)
    document = processor.parse_bytes(pdf.tobytes(), "cv.pdf")
    result = extract_cv(document)
    assert result.name.value == "Casey Placeholder"
    assert [s.text for s in result.skills.value] == ["Python", "SQL", "Tableau"]
    assert result.skills.value[0].source.page_number == 1


def test_unreadable_cv_is_not_extracted(parsed):
    result = extract_cv(parsed["sample_scanned.pdf"])
    assert result.status is ResultStatus.ERROR and result.errors[0].code == "document_ocr_required"
    assert all(field.status is FieldStatus.NOT_EXTRACTED for field in result.fields().values())
    assert CvExtraction.model_validate_json(result.model_dump_json()).model_dump(mode="json") == result.model_dump(mode="json")
