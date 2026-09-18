"""DOCX: headings, paragraphs, list items, tables, headers/footers, reading order; empty and broken files."""

from __future__ import annotations

import io
import zipfile

from docx import Document

from sanad.models.documents import DocumentStatus, ExtractionMethod, FileType, SectionType
from tests.fixtures.documents.builders import empty_docx, image_bytes
from tests.fixtures.documents.synthetic_content import CONTRACT_AR, CONTRACT_EN, CV_EN


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_headings_and_paragraphs_keep_structure_and_order(parsed):
    document = parsed["sample_contract_en.docx"]
    assert document.status is DocumentStatus.SUCCESS
    assert (document.file_type, document.language, document.page_count) == (FileType.DOCX, "en", None)
    title = document.sections[0]
    assert (title.section_type, title.text, title.heading_level, title.style) == (
        SectionType.HEADING, "Employment Contract", 0, "Title")
    probation_heading = next(s for s in document.sections if s.text == "Probation Period")
    assert probation_heading.section_type is SectionType.HEADING and probation_heading.heading_level == 1
    clause = document.sections[document.sections.index(probation_heading) + 1]
    assert clause.section_type is SectionType.PARAGRAPH
    assert clause.title == "Probation Period" and clause.heading_id == probation_heading.section_id
    assert all(s.page_number is None and s.extraction_method is ExtractionMethod.DOCX_XML for s in document.sections)
    orders = [block.order for block in document.blocks()]
    assert orders == sorted(orders) and len(set(orders)) == len(orders)


def test_docx_tables_are_extracted_as_rows_under_their_heading(parsed):
    document = parsed["sample_contract_en.docx"]
    [table] = document.tables
    assert table.rows == [list(row) for row in CONTRACT_EN["compensation_rows"]]
    assert table.title == "Compensation"
    heading = next(s for s in document.sections if s.text == "Compensation")
    assert heading.order < table.order < next(s for s in document.sections if s.text == "Probation Period").order


def test_arabic_docx_text_is_preserved_exactly_with_direction(parsed):
    document = parsed["sample_contract_ar.docx"]
    assert document.status is DocumentStatus.SUCCESS and document.language == "ar"
    assert all(s.text_direction == "rtl" for s in document.sections)
    for _, text in CONTRACT_AR["clauses"]:
        for line in text.split("\n"):
            assert line in document.full_text  # tanween, hamza forms and taa marbuta unchanged
    assert "٢٥٠٠ ريال" in document.full_text  # Arabic-Indic digits are not converted
    assert document.tables[0].rows[0] == ["الراتب الأساسي", "9,500 ريال سعودي شهرياً"]


def test_docx_list_items_are_identified(parsed):
    document = parsed["sample_cv_en.docx"]
    items = [s.text for s in document.sections if s.section_type is SectionType.LIST_ITEM]
    assert items[:4] == CV_EN["skills"]


def test_empty_docx_is_empty_document(processor):
    document = processor.parse_bytes(empty_docx(), "empty.docx")
    assert document.status is DocumentStatus.EMPTY_DOCUMENT
    assert document.sections == [] and document.full_text == ""
    assert any("no text" in w for w in document.extraction_warnings)


def test_docx_with_only_an_image_needs_ocr(processor, tmp_path):
    image = tmp_path / "scan.png"
    image.write_bytes(image_bytes("PNG"))
    document = Document()
    document.add_picture(str(image))
    parsed_document = processor.parse_bytes(_save(document), "scan.docx")
    assert parsed_document.status is DocumentStatus.OCR_REQUIRED
    assert parsed_document.full_text == "" and parsed_document.ocr.required is True


def test_docx_header_footer_nested_and_merged_tables(processor):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Example Tech Solutions LLC - Confidential"
    document.sections[0].footer.paragraphs[0].text = "Page footer text"
    document.add_paragraph("Body paragraph")
    table = document.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Allowances"
    table.cell(0, 2).text = "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text, table.cell(1, 2).text = "Housing", "Monthly", "3,000 SAR"
    inner = table.cell(1, 2).add_table(rows=1, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "Nested", "Value"

    result = processor.parse_bytes(_save(document), "structured.docx")
    assert result.status is DocumentStatus.SUCCESS
    assert result.sections[0].section_type is SectionType.HEADER
    assert result.sections[-1].section_type is SectionType.FOOTER and result.sections[-1].text == "Page footer text"
    outer, nested = result.tables
    assert outer.rows[0] == ["Allowances", "Amount"]  # merged cell not duplicated
    assert nested.parent_table_id == outer.table_id and nested.rows == [["Nested", "Value"]]


def test_malformed_docx_returns_structured_extraction_error(processor, sample_bytes):
    with zipfile.ZipFile(io.BytesIO(sample_bytes["sample_contract_en.docx"])) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries["word/document.xml"] = b"<w:document this is not xml"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for name, data in entries.items():
            target.writestr(name, data)
    document = processor.parse_bytes(buffer.getvalue(), "broken.docx")
    assert document.status is DocumentStatus.EXTRACTION_FAILED
    assert [e.code for e in document.errors] == ["malformed_docx"]
    assert "Traceback" not in document.model_dump_json()


def test_docx_external_entities_are_not_resolved(processor, sample_bytes, tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("TOP-SECRET-LOCAL-FILE-CONTENT", encoding="utf-8")
    with zipfile.ZipFile(io.BytesIO(sample_bytes["sample_contract_en.docx"])) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    xml = entries["word/document.xml"].decode("utf-8")
    header, body = xml.split("?>", 1)
    doctype = f'<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "file://{secret}">]>'
    entries["word/document.xml"] = (header + "?>" + doctype + body.replace("Jordan Sample", "&xxe;", 1)).encode("utf-8")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for name, data in entries.items():
            target.writestr(name, data)
    document = processor.parse_bytes(buffer.getvalue(), "xxe.docx")
    assert "TOP-SECRET-LOCAL-FILE-CONTENT" not in document.model_dump_json()
