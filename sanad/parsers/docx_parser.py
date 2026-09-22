"""DOCX parsing with python-docx, preserving headings, paragraphs, list items, tables and reading order.

python-docx parses the package XML without resolving external entities and never runs macros
(macro-enabled packages are rejected earlier, in detection.py). DOCX has no reliable page
boundaries, so page_number is None for every block.
"""

from __future__ import annotations

import io
import re

from models.documents import DocumentStatus, ExtractionMethod, FileType, ParsedDocument, SectionType
from parsers.base import BaseDocumentParser, DocumentBuilder
from parsers.detection import DocumentSource
from parsers.errors import DocumentErrorCode, DocumentParsingError

_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)
_BULLET_PREFIX = ("•", "●", "▪", "◦", "‣", "- ", "– ")
METHOD = ExtractionMethod.DOCX_XML


def _load_python_docx():
    try:
        import docx  # python-docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentParsingError(DocumentStatus.EXTRACTION_FAILED, DocumentErrorCode.DEPENDENCY_MISSING,
                                   "DOCX support requires the python-docx package.", cause=exc) from None
    return docx, qn, Table, Paragraph


class DocxParser(BaseDocumentParser):
    name = "docx/python-docx"
    file_types = (FileType.DOCX,)

    def parse(self, source: DocumentSource) -> ParsedDocument:
        docx, qn, Table, Paragraph = _load_python_docx()
        try:
            document = docx.Document(io.BytesIO(source.data))
        except Exception as exc:
            raise DocumentParsingError(DocumentStatus.EXTRACTION_FAILED, DocumentErrorCode.MALFORMED_DOCX,
                                       f"'{source.filename}' could not be read as a Word document.", cause=exc) from None

        self._qn, self._Table, self._Paragraph = qn, Table, Paragraph
        self._images = 0
        builder = DocumentBuilder(source, self.name)
        self._add_header_footer(document, builder, "header")
        self._add_container(document.element.body, document, builder)
        self._add_header_footer(document, builder, "footer")

        if builder.sections or builder.tables:
            return builder.build(DocumentStatus.SUCCESS)
        if self._images:
            builder.ocr.required = True
            builder.warn("The document contains images but no text; text inside embedded images was not extracted "
                         "(Sanad's MVP performs no OCR).")
            return builder.build(DocumentStatus.OCR_REQUIRED)
        builder.warn("The document contains no text.")
        return builder.build(DocumentStatus.EMPTY_DOCUMENT)

    # ------------------------------------------------------------------ structure
    def _add_container(self, element, parent, builder: DocumentBuilder, parent_table_id: str | None = None) -> None:
        qn = self._qn
        for child in element.iterchildren():
            if child.tag == qn("w:p"):
                self._add_paragraph(self._Paragraph(child, parent), builder)
            elif child.tag == qn("w:tbl"):
                self._add_table(self._Table(child, parent), builder, parent_table_id)
            elif child.tag == qn("w:sdt"):  # content controls (common in contract templates)
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    self._add_container(content, parent, builder, parent_table_id)

    def _add_paragraph(self, paragraph, builder: DocumentBuilder, section_type: SectionType | None = None) -> None:
        qn = self._qn
        element = paragraph._p
        self._images += len(element.findall(".//" + qn("w:drawing"))) + len(element.findall(".//" + qn("w:pict")))
        text = paragraph.text
        if not text.strip():
            return
        style = self._style_name(paragraph)
        properties = element.pPr
        direction = None
        level = None
        is_list = False
        if properties is not None:
            if properties.find(qn("w:bidi")) is not None:
                direction = "rtl"
            outline = properties.find(qn("w:outlineLvl"))
            if outline is not None and (outline.get(qn("w:val")) or "").isdigit():
                level = int(outline.get(qn("w:val"))) + 1
            is_list = properties.find(qn("w:numPr")) is not None
        if style:
            match = _HEADING_STYLE.match(style.strip())
            if match:
                level = int(match.group(1))
            elif style.strip().lower() == "title":
                level = 0
            elif "list" in style.lower():
                is_list = True

        if section_type is None:
            if level is not None:
                section_type = SectionType.HEADING
            elif is_list or text.lstrip().startswith(_BULLET_PREFIX):
                section_type = SectionType.LIST_ITEM
            elif self._is_bold_heading(paragraph, text):
                section_type = SectionType.HEADING
            else:
                section_type = SectionType.PARAGRAPH
        builder.add_section(text, section_type, METHOD, heading_level=level, style=style, text_direction=direction)

    @staticmethod
    def _style_name(paragraph) -> str | None:
        try:
            return paragraph.style.name if paragraph.style is not None else None
        except (KeyError, AttributeError, ValueError):
            return None

    @staticmethod
    def _is_bold_heading(paragraph, text: str) -> bool:
        stripped = text.strip()
        if len(stripped) > 80 or "\n" in stripped or stripped.endswith((".", "،", ",")):
            return False
        runs = [run for run in paragraph.runs if run.text.strip()]
        return bool(runs) and all(run.bold is True for run in runs)

    def _add_table(self, table, builder: DocumentBuilder, parent_table_id: str | None) -> None:
        rows: list[list[str]] = []
        nested = []
        for row in table.rows:
            cells: list[str] = []
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # horizontally merged cells repeat the same cell
                    continue
                seen.add(id(cell._tc))
                for paragraph in cell.paragraphs:
                    self._images += len(paragraph._p.findall(".//" + self._qn("w:drawing")))
                cells.append("\n".join(paragraph.text for paragraph in cell.paragraphs))
                nested.extend(cell.tables)
            rows.append(cells)
        added = builder.add_table(rows, METHOD, parent_table_id=parent_table_id)
        parent_id = added.table_id if added is not None else parent_table_id
        for inner in nested:
            self._add_table(inner, builder, parent_id)

    def _add_header_footer(self, document, builder: DocumentBuilder, kind: str) -> None:
        seen: set[str] = set()
        section_type = SectionType.HEADER if kind == "header" else SectionType.FOOTER
        for section in document.sections:
            for attribute in (kind, f"first_page_{kind}", f"even_page_{kind}"):
                try:
                    part = getattr(section, attribute)
                    if part.is_linked_to_previous:
                        continue
                    paragraphs = list(part.paragraphs)
                except Exception:
                    builder.warn(f"A document {kind} could not be read.")
                    continue
                for paragraph in paragraphs:
                    text = paragraph.text.strip()
                    if text and text not in seen:
                        seen.add(text)
                        self._add_paragraph(paragraph, builder, section_type=section_type)
